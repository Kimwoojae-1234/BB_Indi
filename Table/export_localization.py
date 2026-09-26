"""Export Public values from the localization master. Python 3.10+, openpyxl.

Generated game data must never be edited instead of the workbook.
Run with --check to validate only, or --check-generated to detect stale JSON.
"""
from __future__ import annotations
import argparse
import json
import os
from pathlib import Path
import re
import string
import tempfile
import sys
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
LANGUAGES = ('Eng', 'Kor', 'Jpn', 'ChnT', 'Esp', 'ChnS', 'Fra', 'Deu', 'Idn', 'Ita', 'Prt', 'Rus', 'Tha', 'Tur', 'Vnm')
DEFAULT_INPUT = ROOT / 'Table/Localization_master_v2.xlsx'
DEFAULT_OUTPUT = ROOT / 'Assets/Resources/Localization/LocalizationItem.json'

def normalize(value):
    if value is None: return ''
    if not isinstance(value, str):
        raise ValueError('Text cells must use the Text format and contain strings.')
    return value.replace('\\n', '\n').replace('\\t', '\t').replace('\r\n', '\n').replace('\r', '\n')

def placeholders(text):
    # .NET composite formatting uses numeric fields (including alignment and format).
    result = set()
    for _, field, _, _ in string.Formatter().parse(text):
        if field is None: continue
        if not re.fullmatch(r'\d+(?:,-?\d+)?', field):
            raise ValueError(f'Invalid composite format field {{{field}}}')
        result.add(int(field.split(',')[0]))
    return result

def read_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if 'Public' not in workbook.sheetnames: raise ValueError('Public sheet is required.')
        sheet = workbook['Public']
        rows = sheet.iter_rows()
        header = [cell.value for cell in next(rows)]
        required = ['key', *LANGUAGES, 'AllowEmpty']
        if header[:6] != ['key', 'Eng', 'Kor', 'Jpn', 'ChnT', 'Esp']:
            raise ValueError('A1:F1 must be key, Eng, Kor, Jpn, ChnT, Esp.')
        if any(header.count(name) != 1 for name in required):
            raise ValueError('Required columns are missing or duplicated: ' + ', '.join(required))
        seen = set(); output = []; warnings = []
        for number, cells in enumerate(rows, 2):
            if all(cell.value is None for cell in cells): continue
            values = dict(zip(header, cells))
            if any(values[name].data_type == 'f' for name in required):
                raise ValueError(f'Row {number}: formulas are not allowed in exported columns.')
            key = normalize(values['key'].value)
            if re.fullmatch(r'TXT_[0-9A-Fa-f]{12}', key):
                raise ValueError(f'Row {number}: use a descriptive key instead of an opaque TXT hash.')
            if not key or key != key.strip() or (re.search(r'\s', key) and key != 'Ingame.DidNo Swing'):
                raise ValueError(f'Row {number}: key is empty or contains whitespace.')
            if key in seen: raise ValueError(f'Row {number}: duplicate key {key}')
            seen.add(key)
            item = {'key': key, **{language: normalize(values[language].value) for language in LANGUAGES}}
            empty_flag = values['AllowEmpty'].value
            if empty_flag not in ('TRUE', 'FALSE', True, False):
                raise ValueError(f'{key}: AllowEmpty must be TRUE or FALSE.')
            allow_empty = empty_flag in ('TRUE', True)
            if allow_empty:
                if any(item[language] for language in LANGUAGES):
                    raise ValueError(f'{key}: AllowEmpty is reserved for intentionally empty entries.')
                item['allowEmpty'] = True
            elif not item['Eng'].strip(): raise ValueError(f'{key}: English is required.')
            if re.search('[가-힣]', item['Eng']): raise ValueError(f'{key}: English contains Korean text.')
            try: expected = placeholders(item['Eng'])
            except ValueError as error: raise ValueError(f'{key}/Eng: {error}') from error
            for language in LANGUAGES[1:]:
                if not item[language]: continue
                try: actual = placeholders(item[language])
                except ValueError as error: raise ValueError(f'{key}/{language}: {error}') from error
                if actual != expected:
                    # Imported translations with incompatible placeholders are kept in the master,
                    # but never shipped. Runtime falls back to English until corrected.
                    warnings.append(f'{key}/{language}: placeholders {sorted(actual)} != {sorted(expected)}; exported blank for English fallback')
                    item[language] = ''
            output.append(item)
        return output, warnings
    finally: workbook.close()

def encode(rows):
    return (json.dumps({'LocalizationItem': rows}, ensure_ascii=False, indent=2) + '\n').encode('utf-8')

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, default=DEFAULT_INPUT)
    parser.add_argument('--output', type=Path, default=DEFAULT_OUTPUT)
    modes = parser.add_mutually_exclusive_group()
    modes.add_argument('--check', action='store_true')
    modes.add_argument('--check-generated', action='store_true')
    args = parser.parse_args()
    try:
        rows, warnings = read_rows(args.input)
        data = encode(rows)
        for warning in warnings: print('WARNING: ' + warning, file=sys.stderr)
        if args.check_generated:
            if not args.output.exists() or args.output.read_bytes() != data:
                raise ValueError(f'Generated data is stale. Run {Path(__file__).name}.')
        elif not args.check:
            args.output.parent.mkdir(parents=True, exist_ok=True)
            fd, temporary = tempfile.mkstemp(dir=args.output.parent, prefix='.localization-', suffix='.tmp')
            try:
                with os.fdopen(fd, 'wb') as stream: stream.write(data)
                os.replace(temporary, args.output)
            finally:
                if os.path.exists(temporary): os.unlink(temporary)
        print(f'PASS: {len(rows)} unique keys; {len(warnings)} preserved translation issues; {args.output}')
        return 0
    except (ValueError, OSError, StopIteration) as error:
        print(f'ERROR: {error}', file=sys.stderr)
        return 1

if __name__ == '__main__': raise SystemExit(main())
